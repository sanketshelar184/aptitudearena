import sys
from pathlib import Path
import pypdf
import docx
from bs4 import BeautifulSoup
from PIL import Image
import io

base = Path(r"d:\Aptitude\Aptitude Pdf's")
all_files = sorted([f for f in base.rglob("*") if f.is_file()])
print(f"Total files found in Aptitude Pdf's: {len(all_files)}", flush=True)

report = []

for idx, f in enumerate(all_files, 1):
    rel = f.relative_to(base)
    ext = f.suffix.lower()
    company = rel.parts[0] if len(rel.parts) > 1 else "Unknown"
    
    info = {
        "index": idx,
        "company": company,
        "filename": f.name,
        "rel_path": str(rel),
        "ext": ext,
        "size_kb": round(f.stat().st_size / 1024, 1),
        "pages": 0,
        "images": 0,
        "image_details": [],
        "text_chars": 0
    }
    
    if ext == ".pdf":
        try:
            reader = pypdf.PdfReader(str(f))
            info["pages"] = len(reader.pages)
            for p_idx, p in enumerate(reader.pages):
                if len(p.images) > 0:
                    info["images"] += len(p.images)
                    for img in p.images:
                        try:
                            pil_img = Image.open(io.BytesIO(img.data))
                            w, h = pil_img.size
                            if len(info["image_details"]) < 10:  # store first 10 for inspection
                                info["image_details"].append(f"p{p_idx+1}:{img.name}({w}x{h})")
                        except Exception:
                            pass
            # sample first 3 pages text
            sample_text = "".join([p.extract_text() or "" for p in reader.pages[:3]])
            info["text_chars"] = len(sample_text)
            print(f"[{idx}/{len(all_files)}] [PDF] {company} / {f.name}: {info['pages']} pgs, {info['images']} imgs, sample_text: {info['text_chars']} chars", flush=True)
        except Exception as e:
            print(f"[{idx}/{len(all_files)}] [PDF ERR] {company} / {f.name}: {e}", flush=True)
            info["error"] = str(e)
            
    elif ext == ".docx":
        try:
            doc = docx.Document(str(f))
            text = "\n".join([p.text for p in doc.paragraphs])
            info["text_chars"] = len(text)
            # check docx inline images
            inline_shapes = 0
            for rel_part in doc.part.related_parts.values():
                if "image" in rel_part.content_type:
                    inline_shapes += 1
            info["images"] = inline_shapes
            print(f"[{idx}/{len(all_files)}] [DOCX] {company} / {f.name}: {len(doc.paragraphs)} paras, {inline_shapes} imgs, {len(text)} chars", flush=True)
        except Exception as e:
            print(f"[{idx}/{len(all_files)}] [DOCX ERR] {company} / {f.name}: {e}", flush=True)
            info["error"] = str(e)

    elif ext in (".htm", ".html"):
        try:
            with open(f, "r", encoding="utf-8", errors="ignore") as fp:
                soup = BeautifulSoup(fp.read(), "html.parser")
                text = soup.get_text()
                imgs = soup.find_all("img")
                info["text_chars"] = len(text)
                info["images"] = len(imgs)
                print(f"[{idx}/{len(all_files)}] [HTM] {company} / {f.name}: {len(text)} chars, {len(imgs)} img tags", flush=True)
        except Exception as e:
            print(f"[{idx}/{len(all_files)}] [HTM ERR] {company} / {f.name}: {e}", flush=True)
            info["error"] = str(e)

    elif ext in (".txt", ".doc"):
        try:
            with open(f, "r", encoding="utf-8", errors="ignore") as fp:
                content = fp.read()
                info["text_chars"] = len(content)
                print(f"[{idx}/{len(all_files)}] [{ext.upper()}] {company} / {f.name}: {len(content)} chars", flush=True)
        except Exception as e:
            print(f"[{idx}/{len(all_files)}] [{ext.upper()} ERR] {company} / {f.name}: {e}", flush=True)
            info["error"] = str(e)
            
    report.append(info)

print("\n=== SUMMARY OF ALL FILES WITH IMAGES ===", flush=True)
for item in report:
    if item.get("images", 0) > 0:
        print(f"COMPANY: {item['company']} | FILE: {item['filename']} | IMAGES: {item['images']} | DETAILS: {item['image_details'][:5]}", flush=True)

